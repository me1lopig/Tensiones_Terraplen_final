
# Libreria de funciones usadas en los cálculos

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Cm, Pt
import os
import io
from datetime import datetime


# ---------------------------------------------------------------------------
# Funciones de cálculo (sin dependencia de archivos Excel)
# ---------------------------------------------------------------------------

def parametro_terreno(cotas, zt):
    """Devuelve el índice de la capa del terreno para una profundidad zt."""
    for z in np.arange(len(cotas) - 1):
        if zt >= cotas[z] and zt <= cotas[z + 1]:
            break
    return z + 1


def n_freatico(nivel_freatico, z):
    if z >= nivel_freatico:
        return (z - nivel_freatico)
    return 0


def insertar_valor(lista, valor):
    lista_mod = lista.copy()
    lista_mod.append(valor)
    lista_mod.sort()
    return lista_mod


def obtener_maximo_menor(lista, valor):
    return max(filter(lambda x: x <= valor, lista))


def presion_total(cotas, valor_nf, pe_saturado, pe_seco, valor_cota):
    lista_cotas = cotas.copy()
    if valor_nf not in lista_cotas:
        lista_valores = insertar_valor(lista_cotas, valor_nf)
    else:
        lista_valores = lista_cotas

    resultado = obtener_maximo_menor(lista_valores, valor_cota)

    peso_saturado = pe_saturado[parametro_terreno(cotas, valor_cota)]
    peso_seco = pe_seco[parametro_terreno(cotas, valor_cota)]

    peso = peso_seco if valor_cota <= valor_nf else peso_saturado
    presion = (valor_cota - resultado) * peso

    for j in range(lista_valores.index(resultado), 0, -1):
        espesor = lista_valores[j] - lista_valores[j - 1]
        ps = pe_saturado[parametro_terreno(cotas, lista_valores[j])]
        pd = pe_seco[parametro_terreno(cotas, lista_valores[j])]
        posicion = lista_valores[j]
        peso = pd if posicion <= valor_nf else ps
        presion += espesor * peso

    return presion


def resistencia_MC(cotas, valor_presion, cohesion, fi, z):
    c = cohesion[parametro_terreno(cotas, z)]
    phi = fi[parametro_terreno(cotas, z)]
    return c + valor_presion * np.tan(np.deg2rad(phi))


def tension_terraplen(a, b, q, x, z):
    beta_a = np.arctan((b - x) / z) + np.arctan((x - a) / z)
    beta_b = np.arctan((x - b) / z) + np.arctan((2 * b - x - a) / z)
    alfa_a = np.arctan((a - x) / z) + np.arctan(x / z)
    alfa_b = np.arctan((a - 2 * b + x) / z) + np.arctan((2 * b - x) / z)

    r_oa = np.sqrt(np.power(x, 2) + np.power(z, 2))
    r_ob = np.sqrt(np.power(2 * b - x, 2) + np.power(z, 2))
    r_1a = np.sqrt(np.power(x - a, 2) + np.power(z, 2))
    r_1b = np.sqrt(np.power(2 * b - x - a, 2) + np.power(z, 2))
    r_22 = np.power(b - x, 2) + np.power(z, 2)

    tensionz = (q / np.pi) * (
        (beta_a + x * alfa_a / a - z * (x - b) / r_22) +
        (beta_b + (2 * b - x) * alfa_b / a - z * (b - x) / r_22)
    )
    tensionx = (q / np.pi) * (
        (beta_a + x * alfa_a / a + z * (x - b) / r_22 + 2 * z * np.log(r_1a / r_oa) / a) +
        beta_b + alfa_b * (2 * b - x) / a + z * (b - x) / r_22 + 2 * z * np.log(r_1b / r_ob) / a
    )
    tensionxz = -(q / np.pi) * (z * (alfa_a + alfa_b) / a - 2 * np.power(z, 2) / r_22)

    return tensionz, tensionx, tensionxz


def asiento_elastico(cotas, z, hi, E, poisson, tensionx, tensionz):
    idx = parametro_terreno(cotas, z)
    return -hi * (tensionz - poisson[idx] * tensionx) / E[idx]


# ---------------------------------------------------------------------------
# Función principal de cálculo (orquesta todo el proceso)
# ---------------------------------------------------------------------------

def calcular_todo(a, b, h, q, ax, incrx, incrz,
                  cotas, nivel_freatico, pe_seco, pe_saturado,
                  E, poisson, cohesion, fi, tipo_calculo):
    """
    Ejecuta el cálculo completo de tensiones y asientos.
    Devuelve un dict con todos los arrays de resultados y coordenadas.
    """
    az = cotas[-1]
    xcoord = np.arange(-(ax + b), ax + b + incrx, incrx)
    zcoord = np.arange(incrz, az + incrz, incrz)

    tension_z = np.zeros((zcoord.size, xcoord.size))
    tension_x = np.zeros((zcoord.size, xcoord.size))
    tension_xz = np.zeros((zcoord.size, xcoord.size))
    tension_z_terreno = np.zeros((zcoord.size, xcoord.size))
    tension_z_efectiva = np.zeros((zcoord.size, xcoord.size))

    asiento = []
    asiento_parcial = 0

    for xarray, x in enumerate(xcoord):
        asiento_parcial = 0
        for zarray, z in enumerate(zcoord):
            tz, tx, txz = tension_terraplen(a, b, q, x + b, z)

            tz0 = presion_total(cotas, nivel_freatico, pe_saturado, pe_seco, z)
            tzef = tz0 - n_freatico(nivel_freatico, z) * 9.81

            tension_z[zarray, xarray] = tz
            tension_x[zarray, xarray] = tx
            tension_xz[zarray, xarray] = txz
            tension_z_terreno[zarray, xarray] = tz0
            tension_z_efectiva[zarray, xarray] = tzef

            tc = tipo_calculo[parametro_terreno(cotas, z)]
            if tc in ['e', 'E']:
                asiento_parcial += asiento_elastico(cotas, z, incrz, E, poisson, tx, tz)

        asiento.append(asiento_parcial)

    asiento = np.array(asiento)
    asiento_max = float(np.min(asiento))

    return {
        'xcoord': xcoord,
        'zcoord': zcoord,
        'tension_z': tension_z,
        'tension_x': tension_x,
        'tension_xz': tension_xz,
        'tension_z_terreno': tension_z_terreno,
        'tension_z_efectiva': tension_z_efectiva,
        'asiento': asiento,
        'asiento_max': asiento_max,
        'a': a, 'b': b, 'h': h, 'q': q,
    }


# ---------------------------------------------------------------------------
# Funciones de visualización (devuelven figura Matplotlib)
# ---------------------------------------------------------------------------

def figura_tensiones(xcoord, zcoord, tension, titulo, tipo, a, b, h):
    """Genera figura de tensiones (contorno o isolínea). Devuelve fig."""
    X, Z = np.meshgrid(xcoord, zcoord)
    fig, ax = plt.subplots(figsize=(9, 5))

    # perfil del terraplén
    xp = [-b, -(b - a), 0, b - a, b]
    yp = [0, h, h, h, 0]
    ax.plot(xp, yp, 'k-', linewidth=1.5)

    if tipo == 'isolinea':
        ax.plot([min(xcoord), max(xcoord)], [0, 0], 'k-', linewidth=0.8)
        curvas = ax.contour(X, -Z, tension, 10)
        ax.clabel(curvas, inline=True, fmt='%2.1f', fontsize=7)
    else:
        curvas = ax.contourf(X, -Z, tension, 10, cmap='RdYlBu_r')
        fig.colorbar(curvas, ax=ax, label='kN/m²')

    ax.set_aspect('equal', adjustable='box')
    ax.set_xlabel('x [m]')
    ax.set_ylabel('z [m]')
    ax.set_title(f'{titulo} [kN/m²]')
    fig.tight_layout()
    return fig


def figura_asientos(xcoord, asiento, b, a, h):
    """Genera figura del perfil de asientos. Devuelve fig."""
    fig, ax = plt.subplots(figsize=(9, 4))
    ax.plot(xcoord, np.array(asiento) * 100, marker='o', markersize=3,
            color='darkred', linestyle='-', linewidth=1.5)
    ax.set_xlabel('x [m]')
    ax.set_ylabel('Asiento [cm]')
    ax.set_title('Perfil de asientos en superficie')
    ax.grid(True, alpha=0.3)
    fig.tight_layout()
    return fig


def figura_perfil_terraplen(a, b, h):
    """Genera croquis esquemático del terraplén."""
    fig, ax = plt.subplots(figsize=(7, 3))
    xp = [-b, -(b - a), 0, b - a, b]
    yp = [0, h, h, h, 0]
    ax.fill(xp, yp, alpha=0.4, color='sienna')
    ax.plot(xp, yp, 'k-', linewidth=2)
    ax.axhline(0, color='gray', linewidth=1)
    ax.fill_between([-b - 2, b + 2], [0, 0], [-1, -1], alpha=0.15, color='green')

    # cotas
    ax.annotate('', xy=(b, 0), xytext=(b - a, 0),
                arrowprops=dict(arrowstyle='<->', color='navy'))
    ax.text(b - a / 2, -0.3, f'a={a} m', ha='center', fontsize=8, color='navy')
    ax.annotate('', xy=(0, 0), xytext=(b - a, 0),
                arrowprops=dict(arrowstyle='<->', color='darkgreen'))
    ax.text((b - a) / 2, -0.6, f'b={b} m', ha='center', fontsize=8, color='darkgreen')
    ax.annotate('', xy=(b + 0.5, h), xytext=(b + 0.5, 0),
                arrowprops=dict(arrowstyle='<->', color='darkred'))
    ax.text(b + 1.2, h / 2, f'H={h} m', ha='left', fontsize=8, color='darkred')

    ax.set_xlim(-b - 2, b + 3)
    ax.set_ylim(-1.2, h + 1)
    ax.set_aspect('equal')
    ax.set_xlabel('x [m]')
    ax.set_ylabel('z [m]')
    ax.set_title('Geometría del terraplén')
    fig.tight_layout()
    return fig


# ---------------------------------------------------------------------------
# Exportación en memoria
# ---------------------------------------------------------------------------

def generar_excel_resultados(resultados, params_terraplen, capas_data):
    """
    Genera un archivo Excel multi-hoja en memoria (BytesIO).
    capas_data: lista de dicts con los datos de cada capa.
    params_terraplen: dict con a, b, h, pe, q, nivel_freatico, incrx, incrz, ax.
    """
    xcoord = resultados['xcoord']
    zcoord = resultados['zcoord']
    output = io.BytesIO()
    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='2E5396')
    center = Alignment(horizontal='center')

    def escribir_matriz(ws, xcoord, zcoord, matriz, label_z='Prof. (m)'):
        ws.cell(1, 1, label_z).font = Font(bold=True)
        ws.cell(1, 1).fill = header_fill
        ws.cell(1, 1).font = header_font
        for c, x in enumerate(xcoord, start=2):
            cell = ws.cell(1, c, round(float(x), 3))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        for r, z in enumerate(zcoord, start=2):
            ws.cell(r, 1, round(float(z), 3))
            for c, val in enumerate(matriz[r - 2, :], start=2):
                ws.cell(r, c, round(float(val), 4))

    def escribir_vector(ws, xcoord, vector, label='x (m)', label2='Asiento (m)'):
        for c, x in enumerate(xcoord, start=1):
            cell = ws.cell(1, c, round(float(x), 3))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        for c, v in enumerate(vector, start=1):
            ws.cell(2, c, round(float(v), 6))
        ws.cell(1, 1).value = label
        ws.insert_rows(1)
        ws.cell(1, 1, label2).font = Font(bold=True)

    # Hoja 1: Datos de entrada
    ws0 = wb.active
    ws0.title = 'Datos_Entrada'
    ws0.cell(1, 1, 'DATOS DEL TERRAPLÉN').font = Font(bold=True, size=12)
    filas_tp = [
        ('Ancho del derrame a (m)', params_terraplen['a']),
        ('Semiancho b (m)', params_terraplen['b']),
        ('Altura H (m)', params_terraplen['h']),
        ('Peso específico γ (kN/m³)', params_terraplen['pe']),
        ('Carga q = γ·H (kN/m²)', params_terraplen['q']),
        ('Nivel freático (m)', params_terraplen['nivel_freatico']),
        ('Incremento malla Δx (m)', params_terraplen['incrx']),
        ('Incremento malla Δz (m)', params_terraplen['incrz']),
        ('Ancho de banda (m)', params_terraplen['ax']),
    ]
    for i, (k, v) in enumerate(filas_tp, start=2):
        ws0.cell(i, 1, k)
        ws0.cell(i, 2, v)
    ws0.column_dimensions['A'].width = 35
    ws0.column_dimensions['B'].width = 15

    ws0.cell(len(filas_tp) + 4, 1, 'PERFIL DEL TERRENO').font = Font(bold=True, size=12)
    headers_cap = ['Capa', 'Espesor (m)', 'Cota inf. (m)', 'γ seco (kN/m³)',
                   'γ sat. (kN/m³)', 'E (kPa)', 'ν', 'c (kPa)', 'φ (°)', 'Tipo cálculo']
    fila_ini = len(filas_tp) + 5
    for c, h_txt in enumerate(headers_cap, start=1):
        cell = ws0.cell(fila_ini, c, h_txt)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
    cota_acum = 0
    for i, capa in enumerate(capas_data, start=1):
        cota_acum += capa['espesor']
        fila = fila_ini + i
        ws0.cell(fila, 1, i)
        ws0.cell(fila, 2, capa['espesor'])
        ws0.cell(fila, 3, round(cota_acum, 3))
        ws0.cell(fila, 4, capa['pe_seco'])
        ws0.cell(fila, 5, capa['pe_sat'])
        ws0.cell(fila, 6, capa['E'])
        ws0.cell(fila, 7, capa['poisson'])
        ws0.cell(fila, 8, capa['cohesion'])
        ws0.cell(fila, 9, capa['fi'])
        ws0.cell(fila, 10, capa['tipo_calculo'])

    # Hojas de tensiones
    matrices = [
        ('Tensiones_Sz', resultados['tension_z']),
        ('Tensiones_Sx', resultados['tension_x']),
        ('Tensiones_Txz', resultados['tension_xz']),
        ('Tension_Total_z', resultados['tension_z_terreno']),
        ('Tension_Efectiva_z', resultados['tension_z_efectiva']),
    ]
    for nombre, mat in matrices:
        ws = wb.create_sheet(nombre)
        escribir_matriz(ws, xcoord, zcoord, mat)

    # Hoja de asientos
    ws_as = wb.create_sheet('Asientos')
    for c, x in enumerate(xcoord, start=1):
        cell = ws_as.cell(1, c, round(float(x), 3))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
    for c, v in enumerate(resultados['asiento'], start=1):
        ws_as.cell(2, c, round(float(v), 6))
    ws_as.cell(3, 1, f"Asiento máximo: {resultados['asiento_max']:.4f} m").font = Font(bold=True)

    wb.save(output)
    output.seek(0)
    return output


def generar_word_informe(resultados, params_terraplen, capas_data, figuras_bytes):
    """
    Genera informe Word en memoria (BytesIO).
    figuras_bytes: dict {nombre: BytesIO con PNG}
    """
    doc = Document()
    doc.add_heading('Análisis Tenso-Deformacional de Terraplén', 0)
    doc.add_paragraph(f'Fecha: {datetime.now().strftime("%d/%m/%Y %H:%M")}')

    # 1. Datos del terraplén
    doc.add_heading('1. Datos del terraplén', level=1)
    filas = [
        ('Ancho del derrame a', params_terraplen['a'], 'm'),
        ('Semiancho de la base b', params_terraplen['b'], 'm'),
        ('Altura H', params_terraplen['h'], 'm'),
        ('Peso específico del relleno γ', params_terraplen['pe'], 'kN/m³'),
        ('Carga aplicada q = γ·H', params_terraplen['q'], 'kN/m²'),
        ('Nivel freático', params_terraplen['nivel_freatico'], 'm'),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, txt in enumerate(['Parámetro', 'Valor', 'Unidad']):
        hdr[i].text = txt
        hdr[i].paragraphs[0].runs[0].font.bold = True
    for nombre, valor, unidad in filas:
        row = table.add_row().cells
        row[0].text = nombre
        row[1].text = f'{valor:.2f}'
        row[2].text = unidad

    doc.add_paragraph()

    # 2. Perfil del terreno
    doc.add_heading('2. Perfil estratigráfico del terreno', level=1)
    cols = ['Capa', 'Espesor (m)', 'γ seco', 'γ sat.', 'E (kPa)', 'ν', 'c (kPa)', 'φ (°)', 'Tipo']
    table2 = doc.add_table(rows=1, cols=len(cols))
    table2.style = 'Table Grid'
    style = doc.styles['Normal']
    style.font.size = Pt(8)
    hdr2 = table2.rows[0].cells
    for i, txt in enumerate(cols):
        hdr2[i].text = txt
        hdr2[i].paragraphs[0].runs[0].font.bold = True
    for i, capa in enumerate(capas_data, start=1):
        row = table2.add_row().cells
        vals = [str(i), str(capa['espesor']), str(capa['pe_seco']), str(capa['pe_sat']),
                str(capa['E']), str(capa['poisson']), str(capa['cohesion']),
                str(capa['fi']), str(capa['tipo_calculo'])]
        for j, v in enumerate(vals):
            row[j].text = v
        for cell in row:
            cell.width = Cm(1.8)

    doc.add_paragraph()

    # 3. Resultados
    doc.add_heading('3. Resultados', level=1)

    asiento_max = resultados['asiento_max']
    doc.add_heading('3.1 Asientos', level=2)
    doc.add_paragraph(
        f'El asiento máximo bajo el eje del terraplén es de {asiento_max*100:.2f} cm '
        f'({asiento_max:.4f} m).'
    )

    tension_max_z = float(np.max(resultados['tension_z']))
    doc.add_paragraph(
        f'El incremento máximo de tensión vertical (Δσz) es de {tension_max_z:.2f} kN/m².'
    )

    # Imágenes
    doc.add_heading('3.2 Mapas de tensiones y asientos', level=2)
    orden = [
        ('Tensiones verticales Δσz — contorno', 'sz_contorno'),
        ('Tensiones verticales Δσz — isolíneas', 'sz_isolinea'),
        ('Tensiones horizontales Δσx — contorno', 'sx_contorno'),
        ('Tensiones tangenciales Δτxz — contorno', 'txz_contorno'),
        ('Asientos', 'asientos'),
    ]
    for titulo_fig, key in orden:
        if key in figuras_bytes:
            doc.add_heading(titulo_fig, level=3)
            buf = figuras_bytes[key]
            buf.seek(0)
            doc.add_picture(buf, width=Cm(14))

    # 4. Conclusiones
    doc.add_heading('4. Conclusiones', level=1)
    doc.add_paragraph(
        f'Bajo las condiciones analizadas con un terraplén de altura H = {params_terraplen["h"]:.1f} m, '
        f'semiancho b = {params_terraplen["b"]:.1f} m y derrame a = {params_terraplen["a"]:.1f} m, '
        f'el incremento máximo de tensión vertical en el subsuelo alcanza {tension_max_z:.1f} kN/m². '
        f'El asiento máximo calculado bajo el eje del terraplén es de {asiento_max*100:.2f} cm.'
    )

    # 5. Bibliografía
    doc.add_heading('5. Bibliografía', level=1)
    doc.add_paragraph(
        'Holl, D.L. (1941). Shearing Stress and Surface Deflections due to Trapezoidal Loads.',
        style='List Bullet'
    )
    doc.add_paragraph(
        'Poulos, H.G. and Davis, E.H. (1974). Elastic Solutions for Soil and Rock Mechanics. Wiley.',
        style='List Bullet'
    )

    buf_word = io.BytesIO()
    doc.save(buf_word)
    buf_word.seek(0)
    return buf_word
